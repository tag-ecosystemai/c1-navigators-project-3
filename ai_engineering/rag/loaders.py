"""Load policy documents (PDF, DOCX, Markdown, TXT) into a common list of text blocks.

Each block is one line or paragraph, tagged with its page number and, when it is a
heading, its heading level. The chunker uses the headings to keep sections together
so that every retrieved chunk can be cited as "Document > Section, page N".
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}

# "3 Affordability", "3.2 DTI referral band", "Article 5 Creditworthiness", "SECTION 4 - ..."
_NUMBERED_HEADING = re.compile(r"^(\d+(?:\.\d+){0,3})\.?\s+([A-Z][^.]{0,90})$")
_WORD_HEADING = re.compile(r"^(?:Article|Section|Chapter|Part|Title|Annex)\s+[\dIVXLC]+[A-Za-z]?\b.{0,90}$", re.I)


@dataclass
class Block:
    text: str
    page: int | None = None
    heading_level: int | None = None  # None = body text


@dataclass
class LoadedDocument:
    source: str  # path relative to the policy folder, used as the citation key
    title: str
    category: str  # first sub-folder under data/policies, or "general"
    blocks: list[Block] = field(default_factory=list)


def detect_heading_level(line: str) -> int | None:
    """Heuristic heading detection for PDFs and plain text, where no styles exist."""
    s = line.strip()
    if not s or len(s) > 100:
        return None
    m = _NUMBERED_HEADING.match(s)
    if m:
        # a heading line should not look like a full sentence
        if s.endswith((".", ";", ",")) or len(m.group(2).split()) > 12:
            return None
        return m.group(1).count(".") + 1
    if _WORD_HEADING.match(s):
        return 1
    if s.isupper() and 3 <= len(s) <= 80 and any(c.isalpha() for c in s):
        return 1
    return None


def _load_markdown(path: Path) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []

    def flush() -> None:
        if para:
            blocks.append(Block(" ".join(para).strip()))
            para.clear()

    for raw in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        line = raw.strip()
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush()
            blocks.append(Block(m.group(2).strip(), heading_level=len(m.group(1))))
        elif not line:
            flush()
        else:
            para.append(line.lstrip("> ").strip())
    flush()
    return blocks


def _load_text(path: Path) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []

    def flush() -> None:
        if para:
            blocks.append(Block(" ".join(para).strip()))
            para.clear()

    for raw in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        line = raw.strip()
        level = detect_heading_level(line)
        if level:
            flush()
            blocks.append(Block(line, heading_level=level))
        elif not line:
            flush()
        else:
            para.append(line)
    flush()
    return blocks


def _load_pdf(path: Path) -> list[Block]:
    from pypdf import PdfReader

    blocks: list[Block] = []
    reader = PdfReader(str(path))
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        para: list[str] = []
        for raw in text.splitlines():
            line = raw.strip()
            level = detect_heading_level(line)
            if level:
                if para:
                    blocks.append(Block(" ".join(para), page=page_no))
                    para = []
                blocks.append(Block(line, page=page_no, heading_level=level))
            elif not line:
                if para:
                    blocks.append(Block(" ".join(para), page=page_no))
                    para = []
            else:
                # PDF lines are hard-wrapped; join them and fix hyphenation
                if para and para[-1].endswith("-"):
                    para[-1] = para[-1][:-1] + line
                else:
                    para.append(line)
        if para:
            blocks.append(Block(" ".join(para), page=page_no))
    return blocks


def _load_docx(path: Path) -> list[Block]:
    import docx

    blocks: list[Block] = []
    for p in docx.Document(str(path)).paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        level = None
        if style.startswith("heading"):
            digits = "".join(c for c in style if c.isdigit())
            level = int(digits) if digits else 1
        elif style == "title":
            level = 1
        blocks.append(Block(text, heading_level=level))
    return blocks


_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".md": _load_markdown,
    ".markdown": _load_markdown,
    ".txt": _load_text,
}


def load_document(path: Path, root: Path) -> LoadedDocument:
    blocks = _LOADERS[path.suffix.lower()](path)
    rel = path.relative_to(root)
    category = rel.parts[0] if len(rel.parts) > 1 else "general"

    # Some templates use a generic H1 ("Retail Banking Credit Underwriting Policy") shared
    # across every document in the folder, with the distinctive name one level down (e.g.
    # "Document Code: POL-RET-2026-01: General Credit Eligibility Standards"). If the first
    # two blocks are a level-1 heading followed immediately by a level-2 heading that carries
    # a document/rule code, drop the generic H1 and promote the specific one to be the title,
    # so citations lead with what actually distinguishes this document.
    if (
        len(blocks) >= 2
        and blocks[0].heading_level == 1
        and blocks[1].heading_level == 2
        and re.search(r"\bcode\b|POL-[A-Z]{2,4}-\d{2,4}", blocks[1].text, re.I)
    ):
        title = blocks[1].text
        blocks = blocks[1:]
    else:
        first_heading = next((b.text for b in blocks if b.heading_level == 1), None)
        title = first_heading or path.stem.replace("_", " ").replace("-", " ")

    return LoadedDocument(source=rel.as_posix(), title=title, category=category, blocks=blocks)


# Folder-notes files, not policy content — skip these even though they're .md
_IGNORED_NAMES = {"readme.md", "readme.txt", "changelog.md"}


def load_directory(root: Path) -> list[LoadedDocument]:
    if not root.exists():
        raise FileNotFoundError(f"Policy folder not found: {root}")
    docs = []
    for path in sorted(root.rglob("*")):
        if (
            path.is_file()
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
            and not path.name.startswith((".", "~$"))
            and path.name.lower() not in _IGNORED_NAMES
        ):
            doc = load_document(path, root)
            if doc.blocks:
                docs.append(doc)
    return docs
